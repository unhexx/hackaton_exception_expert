import pptx
import sys
import json

def dump_deck(path):
    print(f"==================================================")
    print(f"=== {path} ===")
    print(f"==================================================")
    prs = pptx.Presentation(path)
    print(f"Slide width: {prs.slide_width / 914400} in, height: {prs.slide_height / 914400} in")
    print(f"Total slides: {len(prs.slides)}")
    for i, slide in enumerate(prs.slides):
        print(f"\n--- Slide {i+1} (Layout: {slide.slide_layout.name}) ---")
        for s in slide.shapes:
            shape_type = str(s.shape_type)
            pos = f"x={round(s.left/914400, 2)}, y={round(s.top/914400, 2)}, w={round(s.width/914400, 2)}, h={round(s.height/914400, 2)}"
            if s.has_text_frame:
                txt = s.text_frame.text.strip().replace("\n", " \\n ")
                if txt:
                    print(f"  [Text] ({s.name}, {pos}): {txt}")
            elif s.has_table:
                print(f"  [Table] ({s.name}, {pos}): {len(s.table.rows)}x{len(s.table.columns)}")
                for r_idx, row in enumerate(s.table.rows):
                    row_txt = [c.text.strip().replace("\n", " ") for c in row.cells]
                    print(f"    Row {r_idx}: {' | '.join(row_txt)}")
            elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                print(f"  [Picture] ({s.name}, {pos})")
            elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
                print(f"  [Group] ({s.name}, {pos})")
            else:
                print(f"  [Shape {shape_type}] ({s.name}, {pos})")

if __name__ == "__main__":
    for arg in sys.argv[1:]:
        dump_deck(arg)
